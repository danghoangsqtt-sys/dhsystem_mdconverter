from __future__ import annotations

import io
import unittest
from dataclasses import replace
from unittest.mock import patch

from fastapi.testclient import TestClient
from docx import Document

import backend.src.main as main
import backend.src.services.mark_tini.router as mark_tini_router


class ApiSecurityTests(unittest.TestCase):
    def setUp(self) -> None:
        self.client = TestClient(main.app)
        self.auth = {main.API_TOKEN_HEADER: main.settings.api_token}

    def test_health_requires_session_token(self) -> None:
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 401)

        response = self.client.get("/api/health", headers=self.auth)
        self.assertEqual(response.status_code, 200)

    def test_session_rejects_untrusted_browser_origin(self) -> None:
        response = self.client.get(
            "/api/session",
            headers={"Origin": "https://attacker.example"},
        )
        self.assertEqual(response.status_code, 403)
        self.assertNotEqual(response.headers.get("access-control-allow-origin"), "*")

    def test_session_accepts_local_dev_origin(self) -> None:
        response = self.client.get(
            "/api/session",
            headers={"Origin": "http://127.0.0.1:5173"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["token"], main.settings.api_token)

    def test_download_path_traversal_is_rejected_by_uuid_validation(self) -> None:
        response = self.client.get(
            "/api/download/..%5C..%5CREADME",
            headers=self.auth,
        )
        self.assertIn(response.status_code, {404, 422})
        self.assertNotIn("Mark Tini Editor", response.text)

    def test_unsupported_upload_is_rejected_before_queueing(self) -> None:
        response = self.client.post(
            "/api/jobs",
            headers=self.auth,
            files={"file": ("payload.exe", b"MZ", "application/octet-stream")},
        )
        self.assertEqual(response.status_code, 415)

    def test_upload_content_mismatch_is_rejected_and_cleaned_up(self) -> None:
        """A renamed file (e.g. a .txt saved as .pdf) must be caught by
        content sniffing, not just the extension allowlist - and must not
        leave the mislabeled bytes behind in the upload directory."""
        before = set(main.settings.upload_dir.glob("*"))
        response = self.client.post(
            "/api/jobs",
            headers=self.auth,
            files={"file": ("notes.pdf", b"just plain text, not a pdf", "application/pdf")},
        )
        self.assertEqual(response.status_code, 415)
        self.assertEqual(before, set(main.settings.upload_dir.glob("*")))

    def test_upload_content_mismatch_rejects_swapped_ooxml_type(self) -> None:
        """docx and pptx share the same zip container, so a pptx renamed to
        .docx must still be rejected by the internal-member check, not just
        'is this a valid zip'."""
        import io
        import zipfile

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("ppt/presentation.xml", "<p:presentation/>")
        before = set(main.settings.upload_dir.glob("*"))
        response = self.client.post(
            "/api/jobs",
            headers=self.auth,
            files={
                "file": (
                    "slides.docx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 415)
        self.assertEqual(before, set(main.settings.upload_dir.glob("*")))

    def test_upload_content_matching_extension_is_accepted(self) -> None:
        """A minimal but genuine docx (real zip, required member present)
        must pass content validation and reach the job queue."""
        import io
        import zipfile

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/document.xml", "<w:document/>")
        response = self.client.post(
            "/api/jobs",
            headers=self.auth,
            files={
                "file": (
                    "real.docx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 202)
        # Cancel through the real endpoint (not `job_manager.cancel()`
        # directly) so this stays on TestClient's single portal thread,
        # matching how production always calls cancel from the same
        # event loop as the worker - calling the manager directly from the
        # test's own thread would race the worker across threads in a way
        # that can't happen through the real HTTP surface.
        cancel_response = self.client.delete(
            f"/api/jobs/{response.json()['job_id']}",
            headers=self.auth,
        )
        self.assertEqual(cancel_response.status_code, 200)

    def test_verify_citation_requires_session_token(self) -> None:
        response = self.client.post("/api/verify-citation", json={"text": "some claim"})
        self.assertEqual(response.status_code, 401)

    def test_translate_requires_session_token(self) -> None:
        response = self.client.post("/api/translate", json={"text": "some text"})
        self.assertEqual(response.status_code, 401)

    def test_editable_word_export_requires_session_token(self) -> None:
        response = self.client.post(
            "/api/export/markdown-to-word",
            data={"markdown": "# Secret", "original_filename": "paper.pdf"},
        )
        self.assertEqual(response.status_code, 401)

    def test_image_ocr_requires_session_token(self) -> None:
        response = self.client.post(
            "/api/ocr/jobs",
            files={"files": ("photo.png", b"not trusted", "image/png")},
        )
        self.assertEqual(response.status_code, 401)

    def test_image_ocr_rejects_mismatched_magic_and_cleans_upload(self) -> None:
        before = set(main.settings.upload_dir.glob("*"))
        response = self.client.post(
            "/api/ocr/jobs",
            headers=self.auth,
            files={"files": ("photo.png", b"not a png", "image/png")},
        )
        self.assertEqual(response.status_code, 415)
        self.assertEqual(before, set(main.settings.upload_dir.glob("*")))

    def test_image_ocr_rejects_document_formats(self) -> None:
        response = self.client.post(
            "/api/ocr/jobs",
            headers=self.auth,
            files={"files": ("paper.pdf", b"%PDF-1.4", "application/pdf")},
        )
        self.assertEqual(response.status_code, 415)

    def test_editable_word_export_contains_native_text_and_cleans_output(self) -> None:
        docx_before = set(main.settings.output_dir.glob("*.docx"))
        response = self.client.post(
            "/api/export/markdown-to-word",
            headers=self.auth,
            data={
                "markdown": "# Kết quả Docling\n\nVăn bản có thể chỉnh sửa.\n\n| A | B |\n|---|---|\n| 1 | 2 |",
                "original_filename": "fixture.pdf",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn("fixture-chinh-sua.docx", response.headers["content-disposition"])
        document = Document(io.BytesIO(response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Kết quả Docling", text)
        self.assertIn("Văn bản có thể chỉnh sửa", text)
        self.assertEqual(document.tables[0].cell(1, 0).text, "1")
        self.assertEqual(docx_before, set(main.settings.output_dir.glob("*.docx")))

    def test_root_serves_built_frontend_instead_of_api_message(self) -> None:
        """Regression test: `read_root`'s JSON message used to be registered
        unconditionally, which permanently shadowed the StaticFiles mount at
        the same "/" path (Starlette matches routes in registration order) —
        so `start.bat`'s documented `http://localhost:8088` flow showed a bare
        JSON blob instead of the app UI whenever a frontend build existed.
        This only exercises the regression-relevant branch when `frontend/dist`
        has actually been built (e.g. via `npm run build`), same as production.
        """
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        if main.settings.frontend_dist_dir.is_dir():
            self.assertIn("text/html", response.headers.get("content-type", ""))
            self.assertIn('<div id="root">', response.text)
        else:
            self.assertEqual(response.json(), {"message": "Mark Tini Local API"})

    def test_upload_limit_is_enforced_and_partial_file_is_removed(self) -> None:
        tiny_limit_settings = replace(main.settings, max_upload_bytes=4)
        before = set(main.settings.upload_dir.glob("*"))
        with patch.object(mark_tini_router, "settings", tiny_limit_settings):
            response = self.client.post(
                "/api/jobs",
                headers=self.auth,
                files={"file": ("large.pdf", b"12345", "application/pdf")},
            )
        self.assertEqual(response.status_code, 413)
        self.assertEqual(before, set(main.settings.upload_dir.glob("*")))

    def test_job_capacity_exceeded_returns_503_and_cleans_up_upload(self) -> None:
        before = set(main.settings.upload_dir.glob("*"))
        with patch.object(main.job_manager, "submit", side_effect=main.JobCapacityError("queue full")):
            response = self.client.post(
                "/api/jobs",
                headers=self.auth,
                files={"file": ("paper.pdf", b"%PDF-1.4", "application/pdf")},
            )
        self.assertEqual(response.status_code, 503)
        self.assertEqual(before, set(main.settings.upload_dir.glob("*")))


if __name__ == "__main__":
    unittest.main()

