from __future__ import annotations

import base64
import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

from backend.src.services.mark_tini import markdown_to_word_service


class MarkdownToWordServiceTests(unittest.TestCase):
    def test_docling_markdown_becomes_native_editable_word_content(self) -> None:
        markdown = """# Báo cáo Docling

Đây là **văn bản chỉnh sửa được** với *định dạng*.

- Mục thứ nhất
- Mục thứ hai

| Cột A | Cột B |
| --- | --- |
| Dữ liệu 1 | Dữ liệu 2 |

```python
print("Tini")
```
"""
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "editable.docx"
            result = markdown_to_word_service.convert_markdown_to_docx(
                markdown,
                output_path,
                document_title="Báo cáo",
            )

            self.assertEqual(result, output_path)
            document = Document(output_path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Báo cáo Docling", text)
            self.assertIn("văn bản chỉnh sửa được", text)
            self.assertIn("Mục thứ nhất", text)
            self.assertIn('print("Tini")', text)
            self.assertEqual(document.tables[0].cell(1, 0).text, "Dữ liệu 1")
            self.assertEqual(document.core_properties.title, "Báo cáo")
            with zipfile.ZipFile(output_path) as archive:
                media = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(media, [])

    def test_embedded_image_is_written_as_real_centered_picture(self) -> None:
        buffer = io.BytesIO()
        Image.new("RGB", (40, 20), color=(200, 30, 30)).save(buffer, format="PNG")
        data_uri = "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")
        markdown = f"""# Tài liệu có sơ đồ

Đoạn mở đầu.

![Sơ đồ khối]({data_uri})

Đoạn kết thúc.
"""
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "with-image.docx"
            markdown_to_word_service.convert_markdown_to_docx(markdown, output_path)

            document = Document(output_path)
            self.assertEqual(len(document.inline_shapes), 1)
            centered = [p for p in document.paragraphs if p.alignment == WD_ALIGN_PARAGRAPH.CENTER]
            self.assertEqual(len(centered), 1)
            with zipfile.ZipFile(output_path) as archive:
                media = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(len(media), 1)

    def test_corrupt_image_data_falls_back_to_alt_text_without_crashing(self) -> None:
        markdown = "![Sơ đồ hỏng](data:image/png;base64,%%%not-valid%%%)\n"
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "broken-image.docx"
            markdown_to_word_service.convert_markdown_to_docx(markdown, output_path)

            document = Document(output_path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("[Sơ đồ hỏng]", text)
            with zipfile.ZipFile(output_path) as archive:
                media = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(media, [])

    def test_empty_markdown_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            with self.assertRaises(markdown_to_word_service.MarkdownToWordError):
                markdown_to_word_service.convert_markdown_to_docx(
                    "   ",
                    Path(temp_dir) / "empty.docx",
                )

    def test_many_answer_lists_each_restart_instead_of_counting_continuously(self) -> None:
        """Regression test: applying python-docx's built-in "List Number"
        style to every ordered list makes Word share one numbering
        definition document-wide, so on an exam with hundreds of questions
        each question's A/B/C/D choices kept counting up from the previous
        question (observed in production as choices numbered 862, 863,
        864... instead of restarting at 1 for each question)."""
        questions = "\n\n".join(
            f"Câu {n}: Nội dung câu hỏi?\n\n1. Phương án một\n2. Phương án hai\n3. Phương án ba"
            for n in range(1, 40)
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "many-lists.docx"
            markdown_to_word_service.convert_markdown_to_docx(questions, output_path)

            document = Document(output_path)
            marker_paragraphs = [p.text for p in document.paragraphs if p.text.startswith(("1.\t", "2.\t", "3.\t"))]
            # Every question's list must restart at "1." - none of the 39
            # questions' first choice should ever show anything else.
            first_choice_markers = marker_paragraphs[0::3]
            self.assertTrue(all(marker.startswith("1.\t") for marker in first_choice_markers))
            self.assertEqual(len(first_choice_markers), 39)

            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml")
            # No paragraph references Word's shared numbering field - markers
            # are plain literal text, so there is no cross-list counter to
            # leak into (python-docx's blank template ships unused <w:num>
            # boilerplate in numbering.xml regardless; what matters is that
            # no paragraph ever activates it via <w:numPr>).
            self.assertNotIn(b"<w:numPr>", document_xml)

    def test_latex_formula_with_underscore_and_caret_survives_intact(self) -> None:
        """Regression test: markdown-it's emphasis rule treats a bare `_`
        as an italic delimiter. Left unprotected, a formula like `$a_i$`
        could have its underscore consumed as emphasis syntax instead of
        surviving as literal LaTeX, corrupting the formula."""
        markdown = "Công thức: $a_i^2 + \\frac{1}{2} * x$ là kết quả.\n\nKhối: $$\\sum_{i=1}^n a_i$$\n"
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "formula.docx"
            markdown_to_word_service.convert_markdown_to_docx(markdown, output_path)

            document = Document(output_path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("$a_i^2 + \\frac{1}{2} * x$", text)
            self.assertIn("$$\\sum_{i=1}^n a_i$$", text)
            self.assertNotIn("⟦", text)
            self.assertNotIn("FORMULA", text)

    def test_bare_currency_dollar_amount_is_not_treated_as_formula(self) -> None:
        markdown = "Chi phí là $50 cho vé thường và $100 cho vé VIP.\n"
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "currency.docx"
            markdown_to_word_service.convert_markdown_to_docx(markdown, output_path)

            document = Document(output_path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("$50 cho vé thường và $100 cho vé VIP", text)


if __name__ == "__main__":
    unittest.main()
