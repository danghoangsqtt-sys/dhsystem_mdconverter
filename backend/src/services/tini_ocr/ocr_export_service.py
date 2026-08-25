"""Export reviewed OCR pages without running recognition again."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Sequence

from PIL import Image, ImageOps


SUPPORTED_OCR_EXPORTS = {"txt", "markdown", "docx-editable", "docx-faithful"}


def _page_text(page: dict[str, object]) -> str:
    return str(page.get("text", "")).replace("\r\n", "\n").strip()


def _write_editable_docx(pages: Sequence[dict[str, object]], output_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    for index, page in enumerate(pages):
        if len(pages) > 1:
            document.add_heading(str(page.get("filename") or f"Trang {index + 1}"), level=2)
        for line in _page_text(page).splitlines() or [""]:
            document.add_paragraph(line)
        if index < len(pages) - 1:
            document.add_page_break()
    document.save(output_path)


def _write_faithful_docx(image_paths: Sequence[Path], output_path: Path) -> None:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)
    available_width = float(section.page_width - section.left_margin - section.right_margin)
    available_height = float(section.page_height - section.top_margin - section.bottom_margin)

    for index, image_path in enumerate(image_paths):
        with Image.open(image_path) as source:
            corrected = ImageOps.exif_transpose(source).convert("RGB")
            width, height = corrected.size
            image_stream = io.BytesIO()
            corrected.save(image_stream, format="PNG", optimize=False)
            image_stream.seek(0)
        scale = min(available_width / width, available_height / height)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = 0
        run = paragraph.add_run()
        run.add_picture(
            image_stream,
            width=int(width * scale),
            height=int(height * scale),
        )
        if index < len(image_paths) - 1:
            document.add_page_break()
    document.save(output_path)


def build_ocr_export(
    export_format: str,
    pages: Sequence[dict[str, object]],
    output_path: Path,
    image_paths: Sequence[Path] = (),
) -> Path:
    if export_format not in SUPPORTED_OCR_EXPORTS:
        raise ValueError("Định dạng xuất OCR không hợp lệ.")
    if not pages:
        raise ValueError("Không có nội dung OCR để xuất.")
    if export_format == "txt":
        output_path.write_text("\n\n".join(_page_text(page) for page in pages), encoding="utf-8")
    elif export_format == "markdown":
        content = "\n\n".join(
            f"## {page.get('filename') or f'Trang {index + 1}'}\n\n{_page_text(page)}"
            for index, page in enumerate(pages)
        )
        output_path.write_text(content, encoding="utf-8")
    elif export_format == "docx-editable":
        _write_editable_docx(pages, output_path)
    else:
        if len(image_paths) != len(pages):
            raise ValueError("Số ảnh không khớp số trang OCR.")
        _write_faithful_docx(image_paths, output_path)
    return output_path
