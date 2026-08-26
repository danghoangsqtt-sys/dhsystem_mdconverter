"""Create an editable Word document from Docling-produced Markdown.

The document conversion pipeline already pays the expensive Docling cost and
stores its structured Markdown in the editor.  Exporting that reviewed content
directly avoids a second ML pass while still preserving native Word paragraphs,
headings, lists, tables and inline emphasis.
"""

from __future__ import annotations

import base64
import binascii
import io
import re
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from markdown_it import MarkdownIt
from markdown_it.token import Token


class MarkdownToWordError(RuntimeError):
    """Raised when reviewed Markdown cannot be exported to DOCX."""


# Docling emits recognized math/physics/chemistry formulas as literal
# `$...$` / `$$...$$` LaTeX text in the Markdown it exports (docling_core's
# markdown serializer). CommonMark has no concept of math spans, so left
# alone, markdown-it's emphasis/code-span rules partially reinterpret LaTeX
# special characters (`_`, `*`, backticks) as formatting and corrupt the
# formula. Extracting these spans into inert placeholders before parsing
# keeps the original LaTeX byte-for-byte intact end to end.
_BLOCK_FORMULA_RE = re.compile(r"\$\$(.+?)\$\$", re.DOTALL)
_INLINE_FORMULA_RE = re.compile(r"(?<!\$)\$([^\n$]+?)\$(?!\$)")
_LATEX_HINT_RE = re.compile(r"[\\^_{}]")
_FORMULA_PLACEHOLDER_RE = re.compile(r"⟦FORMULA\d+⟧")
_FORMULA_RUN_COLOR = RGBColor(0x5B, 0x2C, 0x8F)


def _extract_formulas(markdown: str) -> tuple[str, dict[str, str]]:
    formulas: dict[str, str] = {}

    def _store(latex_with_delimiters: str) -> str:
        placeholder = f"⟦FORMULA{len(formulas)}⟧"
        formulas[placeholder] = latex_with_delimiters
        return placeholder

    def _replace_block(match: re.Match[str]) -> str:
        return _store(match.group(0))

    def _replace_inline(match: re.Match[str]) -> str:
        if not _LATEX_HINT_RE.search(match.group(1)):
            # Bare "$50" style currency mentions carry no LaTeX markup;
            # leave them as plain text instead of risking a false positive
            # that would visually flag them as a formula.
            return match.group(0)
        return _store(match.group(0))

    text = _BLOCK_FORMULA_RE.sub(_replace_block, markdown)
    text = _INLINE_FORMULA_RE.sub(_replace_inline, text)
    return text, formulas


def _restore_formula_placeholders(text: str, formulas: dict[str, str]) -> str:
    if not formulas or "⟦" not in text:
        return text
    return _FORMULA_PLACEHOLDER_RE.sub(lambda m: formulas.get(m.group(0), m.group(0)), text)


def _split_formula_runs(text: str, formulas: dict[str, str]) -> list[tuple[str, bool]]:
    if not formulas or "⟦" not in text:
        return [(text, False)]
    parts: list[tuple[str, bool]] = []
    last_end = 0
    for match in _FORMULA_PLACEHOLDER_RE.finditer(text):
        placeholder = match.group(0)
        if placeholder not in formulas:
            continue
        if match.start() > last_end:
            parts.append((text[last_end : match.start()], False))
        parts.append((formulas[placeholder], True))
        last_end = match.end()
    if last_end < len(text):
        parts.append((text[last_end:], False))
    return parts


def _decode_data_uri(src: str) -> io.BytesIO | None:
    if not src.startswith("data:image/"):
        return None
    _header, _, payload = src.partition(";base64,")
    if not payload:
        return None
    try:
        return io.BytesIO(base64.b64decode(payload, validate=True))
    except (binascii.Error, ValueError):
        return None


def _sole_image_child(inline_token: Token | None) -> Token | None:
    if inline_token is None:
        return None
    children = inline_token.children or []
    if len(children) == 1 and children[0].type == "image":
        return children[0]
    return None


def _add_image_paragraph(document: DocumentType, token: Token, formulas: dict[str, str]) -> None:
    alt_text = _restore_formula_placeholders(token.content or token.attrGet("alt") or "Hình ảnh", formulas)
    image_stream = _decode_data_uri(token.attrGet("src") or "")

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if image_stream is not None:
        try:
            run = paragraph.add_run()
            picture = run.add_picture(image_stream)
            section = document.sections[0]
            max_width = section.page_width - section.left_margin - section.right_margin
            if max_width > 0 and picture.width > max_width:
                scale = max_width / picture.width
                picture.width = int(picture.width * scale)
                picture.height = int(picture.height * scale)
            return
        except Exception:
            # Corrupt/unsupported image bytes: fall through to the text
            # fallback below instead of failing the whole export.
            pass

    run = paragraph.add_run(f"[{alt_text}]")
    run.italic = True


def _plain_text(tokens: Iterable[Token], formulas: dict[str, str]) -> str:
    pieces: list[str] = []
    for token in tokens:
        if token.type in {"text", "code_inline"}:
            pieces.append(token.content)
        elif token.type in {"softbreak", "hardbreak"}:
            pieces.append("\n")
        elif token.type == "image":
            pieces.append(token.content or token.attrGet("alt") or "[Hình ảnh]")
    return _restore_formula_placeholders("".join(pieces), formulas)


def _add_inline(paragraph, token: Token, formulas: dict[str, str]) -> None:
    children = token.children or []
    bold_depth = 0
    italic_depth = 0
    link_target: str | None = None
    for child in children:
        if child.type == "strong_open":
            bold_depth += 1
            continue
        if child.type == "strong_close":
            bold_depth = max(0, bold_depth - 1)
            continue
        if child.type == "em_open":
            italic_depth += 1
            continue
        if child.type == "em_close":
            italic_depth = max(0, italic_depth - 1)
            continue
        if child.type == "link_open":
            link_target = child.attrGet("href")
            continue
        if child.type == "link_close":
            link_target = None
            continue
        if child.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
            continue
        if child.type == "image":
            text = _restore_formula_placeholders(child.content or child.attrGet("alt") or "Hình ảnh", formulas)
            run = paragraph.add_run(f"[{text}]")
            run.italic = True
            continue
        if child.type not in {"text", "code_inline"}:
            continue
        text = child.content
        if link_target and link_target not in text:
            text = f"{text} ({link_target})"
        is_code = child.type == "code_inline"
        for segment, is_formula in _split_formula_runs(text, formulas):
            if not segment:
                continue
            run = paragraph.add_run(segment)
            if is_formula:
                run.font.name = "Consolas"
                run.font.color.rgb = _FORMULA_RUN_COLOR
            else:
                run.bold = bold_depth > 0
                run.italic = italic_depth > 0
                if is_code:
                    run.font.name = "Consolas"


def _table_rows(
    tokens: Sequence[Token], start: int, formulas: dict[str, str]
) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    row: list[str] | None = None
    cell_parts: list[str] | None = None
    index = start
    while index < len(tokens):
        token = tokens[index]
        if token.type == "table_close":
            return rows, index
        if token.type == "tr_open":
            row = []
        elif token.type in {"th_open", "td_open"}:
            cell_parts = []
        elif token.type == "inline" and cell_parts is not None:
            cell_parts.append(_plain_text(token.children or [], formulas))
        elif token.type in {"th_close", "td_close"} and row is not None and cell_parts is not None:
            row.append("".join(cell_parts).strip())
            cell_parts = None
        elif token.type == "tr_close" and row is not None:
            rows.append(row)
            row = None
        index += 1
    return rows, index


def _add_table(document: DocumentType, rows: Sequence[Sequence[str]]) -> None:
    if not rows:
        return
    column_count = max((len(row) for row in rows), default=0)
    if column_count == 0:
        return
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


class _ListContext:
    """Tracks one nesting level of bullet/ordered list rendering.

    Word's built-in "List Number"/"List Bullet" paragraph styles all share
    one document-wide numbering definition unless a fresh `numId` is
    allocated per list, which python-docx's `style=` shortcut never does.
    On a document with hundreds of independent lists (one per exam
    question's answer choices), that made every choice count continuously
    from the top of the document instead of restarting at each question.
    Rendering the marker as literal, manually-computed text sidesteps
    Word's shared numbering engine entirely so each list is self-contained.
    """

    __slots__ = ("ordered", "counter", "marker_pending")

    def __init__(self, ordered: bool, start: int) -> None:
        self.ordered = ordered
        self.counter = start
        self.marker_pending = False


def convert_markdown_to_docx(
    markdown: str,
    output_path: Path,
    *,
    document_title: str | None = None,
) -> Path:
    """Convert reviewed Docling Markdown to an editable DOCX file."""

    if not markdown.strip():
        raise MarkdownToWordError("Nội dung tài liệu trống, không thể tạo DOCX chỉnh sửa được.")

    try:
        markdown, formulas = _extract_formulas(markdown)
        parser = MarkdownIt("commonmark", {"html": False}).enable("table")
        tokens = parser.parse(markdown)
        document = Document()
        document.core_properties.title = document_title or "Tài liệu Mark Tini"
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)

        list_stack: list[_ListContext] = []
        index = 0
        while index < len(tokens):
            token = tokens[index]
            if token.type == "bullet_list_open":
                list_stack.append(_ListContext(ordered=False, start=1))
            elif token.type == "ordered_list_open":
                start = 1
                if token.attrs:
                    try:
                        start = int(token.attrs.get("start", 1))
                    except (TypeError, ValueError):
                        start = 1
                list_stack.append(_ListContext(ordered=True, start=start))
            elif token.type in {"bullet_list_close", "ordered_list_close"}:
                if list_stack:
                    list_stack.pop()
            elif token.type == "list_item_open":
                if list_stack:
                    list_stack[-1].marker_pending = True
            elif token.type == "heading_open":
                level = min(9, max(1, int(token.tag[1:] or "1")))
                paragraph = document.add_heading(level=level)
                if index + 1 < len(tokens) and tokens[index + 1].type == "inline":
                    _add_inline(paragraph, tokens[index + 1], formulas)
            elif token.type == "paragraph_open":
                inline_token = tokens[index + 1] if index + 1 < len(tokens) and tokens[index + 1].type == "inline" else None
                sole_image = _sole_image_child(inline_token)
                if sole_image is not None:
                    _add_image_paragraph(document, sole_image, formulas)
                else:
                    paragraph = document.add_paragraph()
                    if list_stack:
                        depth = len(list_stack) - 1
                        base_indent = Pt(18 + depth * 18)
                        ctx = list_stack[-1]
                        if ctx.marker_pending:
                            ctx.marker_pending = False
                            marker_text = f"{ctx.counter}." if ctx.ordered else "•"
                            if ctx.ordered:
                                ctx.counter += 1
                            paragraph.paragraph_format.left_indent = base_indent
                            paragraph.paragraph_format.first_line_indent = Pt(-18)
                            paragraph.paragraph_format.tab_stops.add_tab_stop(base_indent)
                            paragraph.add_run(f"{marker_text}\t")
                        else:
                            # Continuation paragraph within a multi-paragraph
                            # list item: align under the marker, no new one.
                            paragraph.paragraph_format.left_indent = base_indent
                    if inline_token is not None:
                        _add_inline(paragraph, inline_token, formulas)
            elif token.type in {"fence", "code_block"}:
                paragraph = document.add_paragraph(style="No Spacing")
                run = paragraph.add_run(_restore_formula_placeholders(token.content.rstrip("\n"), formulas))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif token.type == "table_open":
                rows, end_index = _table_rows(tokens, index + 1, formulas)
                _add_table(document, rows)
                index = end_index
            elif token.type == "hr":
                document.add_paragraph("―" * 24)
            index += 1

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        return output_path
    except MarkdownToWordError:
        raise
    except Exception as exc:
        raise MarkdownToWordError("Không thể tạo DOCX chỉnh sửa được từ nội dung Docling.") from exc
