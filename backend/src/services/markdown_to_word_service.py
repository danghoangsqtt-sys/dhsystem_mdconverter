"""Create an editable Word document from Docling-produced Markdown.

The document conversion pipeline already pays the expensive Docling cost and
stores its structured Markdown in the editor.  Exporting that reviewed content
directly avoids a second ML pass while still preserving native Word paragraphs,
headings, lists, tables and inline emphasis.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt
from markdown_it import MarkdownIt
from markdown_it.token import Token


class MarkdownToWordError(RuntimeError):
    """Raised when reviewed Markdown cannot be exported to DOCX."""


def _plain_text(tokens: Iterable[Token]) -> str:
    pieces: list[str] = []
    for token in tokens:
        if token.type in {"text", "code_inline"}:
            pieces.append(token.content)
        elif token.type in {"softbreak", "hardbreak"}:
            pieces.append("\n")
        elif token.type == "image":
            pieces.append(token.content or token.attrGet("alt") or "[Hình ảnh]")
    return "".join(pieces)


def _add_inline(paragraph, token: Token) -> None:
    children = token.children or []
    bold_depth = 0
    italic_depth = 0
    link_target: str | None = None
    for child in children:
        if child.type == "strong_open":
            bold_depth += 1
            continue
        if child.type == "strong_close":
            bold_depth = max(0, bold_depth - 1)
            continue
        if child.type == "em_open":
            italic_depth += 1
            continue
        if child.type == "em_close":
            italic_depth = max(0, italic_depth - 1)
            continue
        if child.type == "link_open":
            link_target = child.attrGet("href")
            continue
        if child.type == "link_close":
            link_target = None
            continue
        if child.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
            continue
        if child.type == "image":
            text = child.content or child.attrGet("alt") or "Hình ảnh"
            run = paragraph.add_run(f"[{text}]")
            run.italic = True
            continue
        if child.type not in {"text", "code_inline"}:
            continue
        text = child.content
        if link_target and link_target not in text:
            text = f"{text} ({link_target})"
        run = paragraph.add_run(text)
        run.bold = bold_depth > 0
        run.italic = italic_depth > 0
        if child.type == "code_inline":
            run.font.name = "Consolas"


def _table_rows(tokens: Sequence[Token], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    row: list[str] | None = None
    cell_parts: list[str] | None = None
    index = start
    while index < len(tokens):
        token = tokens[index]
        if token.type == "table_close":
            return rows, index
        if token.type == "tr_open":
            row = []
        elif token.type in {"th_open", "td_open"}:
            cell_parts = []
        elif token.type == "inline" and cell_parts is not None:
            cell_parts.append(_plain_text(token.children or []))
        elif token.type in {"th_close", "td_close"} and row is not None and cell_parts is not None:
            row.append("".join(cell_parts).strip())
            cell_parts = None
        elif token.type == "tr_close" and row is not None:
            rows.append(row)
            row = None
        index += 1
    return rows, index


def _add_table(document: DocumentType, rows: Sequence[Sequence[str]]) -> None:
    if not rows:
        return
    column_count = max((len(row) for row in rows), default=0)
    if column_count == 0:
        return
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def convert_markdown_to_docx(
    markdown: str,
    output_path: Path,
    *,
    document_title: str | None = None,
) -> Path:
    """Convert reviewed Docling Markdown to an editable DOCX file."""

    if not markdown.strip():
        raise MarkdownToWordError("Nội dung tài liệu trống, không thể tạo DOCX chỉnh sửa được.")

    try:
        parser = MarkdownIt("commonmark", {"html": False}).enable("table")
        tokens = parser.parse(markdown)
        document = Document()
        document.core_properties.title = document_title or "Tài liệu Mark Tini"
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)

        list_stack: list[str] = []
        index = 0
        while index < len(tokens):
            token = tokens[index]
            if token.type == "bullet_list_open":
                list_stack.append("List Bullet")
            elif token.type == "ordered_list_open":
                list_stack.append("List Number")
            elif token.type in {"bullet_list_close", "ordered_list_close"}:
                if list_stack:
                    list_stack.pop()
            elif token.type == "heading_open":
                level = min(9, max(1, int(token.tag[1:] or "1")))
                paragraph = document.add_heading(level=level)
                if index + 1 < len(tokens) and tokens[index + 1].type == "inline":
                    _add_inline(paragraph, tokens[index + 1])
            elif token.type == "paragraph_open":
                style = list_stack[-1] if list_stack else None
                paragraph = document.add_paragraph(style=style)
                if index + 1 < len(tokens) and tokens[index + 1].type == "inline":
                    _add_inline(paragraph, tokens[index + 1])
            elif token.type in {"fence", "code_block"}:
                paragraph = document.add_paragraph(style="No Spacing")
                run = paragraph.add_run(token.content.rstrip("\n"))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif token.type == "table_open":
                rows, end_index = _table_rows(tokens, index + 1)
                _add_table(document, rows)
                index = end_index
            elif token.type == "hr":
                document.add_paragraph("―" * 24)
            index += 1

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        return output_path
    except MarkdownToWordError:
        raise
    except Exception as exc:
        raise MarkdownToWordError("Không thể tạo DOCX chỉnh sửa được từ nội dung Docling.") from exc
